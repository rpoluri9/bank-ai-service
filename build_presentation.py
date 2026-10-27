#!/usr/bin/env python3
"""Build the 10-slide Bank AI Platform presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

BASE = "/home/hatch/workspace/job-search/project-13-bank-ai-service"


def slide(title, bullets, image=None):
    s = prs.slides.add_slide(BLANK)
    tx = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    if image:
        s.shapes.add_picture(image, Inches(0.6), Inches(1.4), width=Inches(6.2))
        bx = Inches(7.1)
        bw = Inches(5.6)
    else:
        bx = Inches(0.6)
        bw = Inches(12.1)
    box = s.shapes.add_textbox(bx, Inches(1.4), bw, Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(22)
        p.space_after = Pt(14)
        p.level = 0
    return s


slide("Bank AI Platform", [
    "Banking + AI assistant, built as a portfolio project",
    "By Ramesh Poluri — Java + Python developer",
    "Spring Boot banking microservice + Python RAG agent",
    "GitHub: github.com/rpoluri9/bank-ai-service",
])

slide("The idea", [
    "A bank customers can talk to: ask about fees, policies, and your own money",
    "Document questions answered from real bank documents (RAG)",
    "Balance and transaction questions answered from live data — never guessed",
    "Built the way a Java developer's project is built: microservice, REST, CI/CD",
])

slide("Architecture", [
    "React frontend on :5173 logs in with the demo user",
    "User / Postman talks to the AI service (:8081)",
    "AI service calls the banking service (:8080) through tools, with Basic auth",
    "Banking service persists accounts + transactions in MySQL (JPA)",
    "Bank documents live in a TF-IDF index for retrieval",
    "Transaction events flow through a publisher interface",
    "Docker + GitHub Actions wrap all services",
], image=f"{BASE}/docs/architecture.png")

slide("Banking microservice — Spring Boot", [
    "Java 17, Spring Boot 3.2, REST APIs",
    "JPA entities + repositories on MySQL (H2 in tests)",
    "Spring Security: HTTP Basic auth on every /api endpoint",
    "React frontend (Vite) on :5173 for accounts and the AI panel",
    "Create account, deposit, withdraw, transfer, history",
    "Request validation + centralized exception handling",
    "Every error returns the same clean JSON contract",
])

slide("Core banking rules", [
    "Balances are BigDecimal — money is never a double",
    "No overdraft: withdrawals and transfers that would go below zero fail",
    "Every movement writes an immutable Transaction (audit trail)",
    "Transfers lock both accounts in a fixed order — no deadlocks",
    "20 JUnit 5 + MockMvc tests, all passing (13 service + 4 security + 3 repository)",
], image=f"{BASE}/docs/transaction-flow.png")

slide("RAG pipeline — pure Python, offline", [
    "Bank documents: fee schedule, policies, FAQs, loan basics",
    "Parse → chunk (max 800 chars) → TF-IDF index",
    "TF-IDF vectors play the role of embeddings — no model download",
    "Top-3 chunks by cosine similarity; weak matches are dropped",
    "Standard library only — no AI libraries, no API keys",
], image=f"{BASE}/docs/rag-pipeline.png")

slide("Agent + tool calling", [
    "ADK-style loop: instructions → plan tools → call tools → answer",
    "get_balance and list_transactions call the banking REST APIs",
    "One question can trigger both a tool call and a document lookup",
    "Vertex AI provider interface included; local provider is the default",
    "If neither tools nor documents have the answer, it says so honestly",
])

slide("Events — Pub/Sub interface, local default", [
    "Every transaction publishes a TransactionEvent",
    "Default: local publisher (in-memory + logged) — runs offline",
    "Configured: real Google Cloud Pub/Sub via its REST API",
    "Pub/Sub only activates with explicit project, topic, and token",
    "Same interface on both services — swap without changing code",
])

slide("Local vs cloud — stated honestly", [
    "Out of the box, nothing talks to Google Cloud — no keys needed",
    "Vertex AI and Pub/Sub are opt-in provider implementations",
    "Verified live on real MySQL: 1000 → 1500 → 1300, overdraft rejected (400), transfer moved 300 between accounts",
    "AI answered a fee question from documents and a balance question via tool call",
    "Tests: banking 20/20, AI service 14/14",
])

slide("Resume points — Bank AI Platform", [
    "Built a banking microservice: accounts, deposits, withdrawals, transfers",
    "Enforced core banking rules: no overdraft, full audit trail per transaction",
    "Built Spring Boot REST APIs with validation and centralized exception handling",
    "Persisted data with Spring Data JPA on MySQL; secured APIs with Spring Security",
    "Built a React frontend for accounts, transfers, and the AI assistant panel",
    "Built a Python AI service with RAG pipelines over bank documents",
    "Added document ingestion, tool calling, Pub/Sub events, Docker, and CI/CD",
])

prs.save("/home/hatch/workspace/your_files/bank-ai-service-presentation.pptx")
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
